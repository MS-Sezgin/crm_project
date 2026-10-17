"""Excel data access for the CRM pages.

Data files live in the data/ folder: Applications.xlsx, Interviews.xlsx,
Mentor.xlsx and Users.xlsx. Columns are found by keywords, so small wording
differences in the headers do not break the app.

Set the CRM_DATA_DIR environment variable to read files from another folder
(the test script uses this).
"""
import os
import unicodedata
from collections import Counter
from datetime import datetime

from openpyxl import load_workbook

ROOT_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.environ.get("CRM_DATA_DIR", os.path.join(ROOT_DIR, "data"))

APPLICATIONS_FILE = "Applications.xlsx"
INTERVIEWS_FILE = "Interviews.xlsx"
MENTOR_FILE = "Mentor.xlsx"
USERS_FILE = "Users.xlsx"
VIT1_FILE = "VIT1.xlsx"
VIT2_FILE = "VIT2.xlsx"

CALENDAR_COLUMNS = ["Activity Name", "Start Time", "Participant Mail", "Organizer Mail"]
NAME_KEYWORDS = ("name", "soyad")

_cache = {}


def normalize(value):
    """Lower-case text without accents, so keyword matching works for Turkish headers."""
    text = "" if value is None else str(value)
    decomposed = unicodedata.normalize("NFKD", text)
    plain = "".join(ch for ch in decomposed if not unicodedata.combining(ch))
    return plain.replace("\u0131", "i").lower().strip()


def find_any(headers, keywords):
    """Index of the first header that contains any of the keywords, or None."""
    for index, header in enumerate(headers):
        name = normalize(header)
        if any(keyword in name for keyword in keywords):
            return index
    return None


def _format(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.hour == 0 and value.minute == 0:
            return value.strftime("%d.%m.%Y")
        return value.strftime("%d.%m.%Y %H:%M")
    return str(value).strip()


def load_sheet(file_name):
    """Return (headers, rows) from the first sheet. Every cell is returned as text."""
    path = os.path.join(DATA_DIR, file_name)
    if not os.path.exists(path):
        raise FileNotFoundError(f"{file_name} was not found in the data folder.")

    modified = os.path.getmtime(path)
    cached = _cache.get(path)
    if cached and cached[0] == modified:
        return cached[1], cached[2]

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        raw = [list(row) for row in workbook.worksheets[0].iter_rows(values_only=True)]
    finally:
        workbook.close()
    raw = [row for row in raw if any(v not in (None, "") for v in row)]

    if not raw:
        headers, rows = [], []
    else:
        width = max(len(row) for row in raw)
        header_row = raw[0]
        headers = []
        for i in range(width):
            title = header_row[i] if i < len(header_row) else None
            headers.append(str(title).strip() if title not in (None, "") else f"Column {i + 1}")
        rows = [[_format(row[i]) if i < len(row) else "" for i in range(width)] for row in raw[1:]]

    _cache[path] = (modified, headers, rows)
    return headers, rows


def search_rows(headers, rows, text):
    """Keep rows where a name starts with the typed text (e.g. 'as' finds 'Asli Kaya')."""
    query = normalize(text)
    if not query:
        return rows
    name_cols = [i for i, h in enumerate(headers) if any(k in normalize(h) for k in NAME_KEYWORDS)]
    if not name_cols:
        name_cols = list(range(len(headers)))
    result = []
    for row in rows:
        for i in name_cols:
            cell = normalize(row[i])
            if cell.startswith(query) or any(word.startswith(query) for word in cell.split()):
                result.append(row)
                break
    return result


def check_login(username, password):
    """Return "admin", "user", or None when the username or password is wrong."""
    headers, rows = load_sheet(USERS_FILE)
    user_col = find_any(headers, ("user", "kullan"))
    pass_col = find_any(headers, ("pass", "sifre"))
    role_col = find_any(headers, ("auth", "yetki", "role"))
    if None in (user_col, pass_col, role_col):
        raise ValueError("Users.xlsx needs username, password and authority columns.")
    wanted = username.strip().lower()
    for row in rows:
        if normalize(row[user_col]) == wanted and row[pass_col] == password:
            return "admin" if "admin" in normalize(row[role_col]) else "user"
    return None


def applications_by_mentor(defined):
    """Applications that have (defined=True) or have not (defined=False) a mentor meeting."""
    headers, rows = load_sheet(APPLICATIONS_FILE)
    col = find_any(headers, ("mentor",))
    if col is None:
        raise ValueError("No mentor meeting column found in Applications.xlsx.")
    return headers, [r for r in rows if (r[col] != "") == defined]


def interviews_by_status(kind):
    """kind: "sent", "received", or None for all rows."""
    headers, rows = load_sheet(INTERVIEWS_FILE)
    if kind is None:
        return headers, rows
    keywords = {"sent": ("gonder", "sent"), "received": ("gelis", "received", "arriv")}[kind]
    col = find_any(headers, keywords)
    if col is None:
        raise ValueError("No project column found in Interviews.xlsx.")
    return headers, [r for r in rows if r[col] != ""]


def mentor_categories():
    """Distinct values of the mentor direction column, used for the tab combo box."""
    headers, rows = load_sheet(MENTOR_FILE)
    col = find_any(headers, ("recommend", "yonlendir", "direct"))
    if col is None:
        return []
    return sorted({r[col] for r in rows if r[col]})


def mentor_by_category(category):
    """Mentor rows for one category, or all rows when category is None."""
    headers, rows = load_sheet(MENTOR_FILE)
    if category is None:
        return headers, rows
    col = find_any(headers, ("recommend", "yonlendir", "direct"))
    if col is None:
        return headers, rows
    return headers, [r for r in rows if r[col] == category]


def _name_mail_cols(headers):
    return find_any(headers, NAME_KEYWORDS), find_any(headers, ("mail",))


def _cell(row, col):
    return row[col] if col is not None else ""


def duplicate_applications():
    """Applications with the same name and e-mail registered more than once."""
    headers, rows = load_sheet(APPLICATIONS_FILE)
    n, m = _name_mail_cols(headers)
    key = lambda r: (normalize(_cell(r, n)), normalize(_cell(r, m)))
    counts = Counter(key(r) for r in rows)
    return headers, [r for r in rows if counts[key(r)] > 1]


def unique_applications():
    """Each name shown once, even if it was registered several times."""
    headers, rows = load_sheet(APPLICATIONS_FILE)
    n, _ = _name_mail_cols(headers)
    seen, result = set(), []
    for r in rows:
        name = normalize(_cell(r, n))
        if name in seen:
            continue
        seen.add(name)
        result.append(r)
    return headers, result


def _people_by_file():
    """Map (name, e-mail) -> {"name", "mail", "files"} across Applications, VIT1 and VIT2."""
    sources = [("Applications", APPLICATIONS_FILE), ("VIT1", VIT1_FILE), ("VIT2", VIT2_FILE)]
    people = {}
    for label, file_name in sources:
        headers, rows = load_sheet(file_name)
        n, m = _name_mail_cols(headers)
        for r in rows:
            name, mail = _cell(r, n), _cell(r, m)
            key = (normalize(name), normalize(mail))
            if not any(key):
                continue
            entry = people.setdefault(key, {"name": name, "mail": mail, "files": []})
            if label not in entry["files"]:
                entry["files"].append(label)
    return list(people.values())


def previous_vit_check():
    """People found in at least two of: Applications, VIT1, VIT2 (applied to more than one VIT)."""
    return [p for p in _people_by_file() if len(p["files"]) >= 2]


def different_records():
    """People found in exactly one of VIT1 and VIT2 (not common to both)."""
    result = []
    for p in _people_by_file():
        in_vits = {f for f in p["files"] if f in ("VIT1", "VIT2")}
        if len(in_vits) == 1:
            result.append(p)
    return result


def people_as_table(people):
    headers = ["Full Name", "Email", "Found In"]
    rows = [[p["name"], p["mail"], ", ".join(p["files"])] for p in people]
    return headers, rows


DRIVE_FILES = [APPLICATIONS_FILE, INTERVIEWS_FILE, MENTOR_FILE, USERS_FILE, VIT1_FILE, VIT2_FILE]


def sync_from_drive():
    """Refresh the local data/ copies from Google Drive. Returns a short status message.

    If Drive cannot be reached, the existing local files are used.
    """
    try:
        import google_services
        results = google_services.sync_drive_files(DATA_DIR, DRIVE_FILES)
    except Exception as error:  # no internet, no credentials, file locked, API error
        return f"Drive sync skipped, using local files: {error}"
    updated = [name for name, state in results.items() if state == "downloaded"]
    if updated:
        return f"Drive: {len(updated)} file(s) updated"
    return "Drive: local files are up to date"

